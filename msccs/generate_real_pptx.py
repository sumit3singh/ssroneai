"""
Generate high-fidelity, light-themed, professional Microsoft PowerPoint (.pptx) files
with slide titles, styled bullet points, university headers, and presenter speech notes.
"""
import os
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
APP_DIR = os.path.join(BASE_DIR, "app")
DOWNLOADS_DIR = os.path.join(APP_DIR, "downloads")
os.makedirs(DOWNLOADS_DIR, exist_ok=True)

# Elegant Light Theme Colors
NAVY = RGBColor(15, 30, 54)        # #0F1E36 - Deep University Navy
LIGHT_BLUE = RGBColor(239, 246, 255) # #EFF6FF - Soft Ice Blue Background
GOLD = RGBColor(197, 139, 58)      # #C58B3A - University Gold
SLATE = RGBColor(30, 41, 59)       # #1E293B - Dark Slate Text
MUTED_SLATE = RGBColor(100, 116, 139) # #64748B - Muted Text
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(255, 255, 255)

def clean_text(text):
    """Strip raw markdown formatting for PowerPoint text."""
    if not text:
        return ""
    t = text.replace("**", "").replace("`", "").replace("###", "").replace("##", "")
    t = t.replace("[Student Name]", "Sumit Singh")
    t = t.replace("[Candidate Name]", "Sumit Singh")
    t = t.replace("[Roll Number]", "PSOU/MSCS/2026/042")
    t = t.replace("[Student Roll Number]", "PSOU/MSCS/2026/042")
    t = t.replace("[Supervisor Name]", "Dr. Harpreet Singh")
    t = t.replace("[Supervisor / Faculty Guide Name]", "Dr. Harpreet Singh")
    t = t.replace("[Jan / July Admission Cycle]", "July 2024 - 2026")
    t = t.replace("[Study Centre Code & Name]", "JGND PSOU Regional Centre, Patiala")
    return t.strip()

def create_presentation(slides_data, file_name, is_major_project=True):
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 Widescreen
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Title Slide (Light, Pristine, University Branded)
    title_slide = prs.slides.add_slide(blank_layout)
    
    # Background card
    bg = title_slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5)) # MSO_SHAPE.RECTANGLE = 1
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BLUE
    bg.line.color.rgb = LIGHT_BLUE

    # University Header Text Box
    u_box = title_slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(0.6))
    u_tf = u_box.text_frame
    u_p = u_tf.paragraphs[0]
    u_p.text = "JAGAT GURU NANAK DEV PUNJAB STATE OPEN UNIVERSITY, PATIALA"
    u_p.font.size = Pt(15)
    u_p.font.bold = True
    u_p.font.color.rgb = GOLD
    u_p.alignment = PP_ALIGN.CENTER

    # Degree & Department
    d_box = title_slide.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(11.333), Inches(0.5))
    d_tf = d_box.text_frame
    d_p = d_tf.paragraphs[0]
    d_p.text = "DEPARTMENT OF COMPUTER SCIENCE & APPLICATIONS | M.SC. (CS) SEMESTER-IV"
    d_p.font.size = Pt(12)
    d_p.font.bold = True
    d_p.font.color.rgb = NAVY
    d_p.alignment = PP_ALIGN.CENTER

    # Main Project Title Box
    t_box = title_slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.933), Inches(2.2))
    t_tf = t_box.text_frame
    t_tf.word_wrap = True
    t_p = t_tf.paragraphs[0]
    t_p.text = clean_text(slides_data[0]["title"] if is_major_project else slides_data[0]["title"])
    if is_major_project:
        t_p.text = "Design and Implementation of a High-Throughput, Offline-First, Multi-Tenant SaaS Enterprise POS Architecture with Integrated AI Copilot and Computer Vision OCR"
    else:
        t_p.text = "Architectural Patterns for Offline-First Distributed Web Applications: Balancing Zero-Latency User Experience with Eventual Consistency"
    t_p.font.size = Pt(24)
    t_p.font.bold = True
    t_p.font.color.rgb = NAVY
    t_p.alignment = PP_ALIGN.CENTER

    # Student Particulars Card (White floating container)
    p_card = title_slide.shapes.add_shape(1, Inches(2.5), Inches(4.7), Inches(8.333), Inches(1.9))
    p_card.fill.solid()
    p_card.fill.fore_color.rgb = WHITE
    p_card.line.color.rgb = RGBColor(226, 232, 240)
    
    p_tf = p_card.text_frame
    p_tf.word_wrap = True
    p1 = p_tf.paragraphs[0]
    p1.text = "Candidate: Sumit Singh   |   Roll Number: PSOU/MSCS/2026/042"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = NAVY
    p1.alignment = PP_ALIGN.CENTER

    p2 = p_tf.add_paragraph()
    p2.text = "Faculty Supervisor: Dr. Harpreet Singh (Assistant Professor)"
    p2.font.size = Pt(12)
    p2.font.color.rgb = SLATE
    p2.alignment = PP_ALIGN.CENTER

    p3 = p_tf.add_paragraph()
    p3.text = "Course: " + ("MSCS-4-03P (Major Project — 200 Marks)" if is_major_project else "MSCS-4-04 (Seminar — 50 Marks)")
    p3.font.size = Pt(11)
    p3.font.bold = True
    p3.font.color.rgb = GOLD
    p3.alignment = PP_ALIGN.CENTER

    # Add Speaker Speech Script to Title Slide Presenter Notes
    if slides_data[0].get("speech"):
        title_slide.notes_slide.notes_text_frame.text = clean_text(slides_data[0]["speech"])

    # Content Slides (Slides 2 to N)
    for idx, slide_item in enumerate(slides_data[1:], start=2):
        slide = prs.slides.add_slide(blank_layout)

        # Background (Light Clean White/Slate-50)
        s_bg = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
        s_bg.fill.solid()
        s_bg.fill.fore_color.rgb = RGBColor(248, 250, 252) # Slate-50
        s_bg.line.color.rgb = RGBColor(248, 250, 252)

        # Top White Header Card with Gold accent
        h_card = slide.shapes.add_shape(1, Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.0))
        h_card.fill.solid()
        h_card.fill.fore_color.rgb = WHITE
        h_card.line.color.rgb = RGBColor(226, 232, 240)

        # Gold accent bar on left of header
        accent = slide.shapes.add_shape(1, Inches(0.8), Inches(0.5), Inches(0.15), Inches(1.0))
        accent.fill.solid()
        accent.fill.fore_color.rgb = GOLD
        accent.line.color.rgb = GOLD

        # Title Text
        h_tf = h_card.text_frame
        h_tf.word_wrap = True
        h_p = h_tf.paragraphs[0]
        h_p.text = clean_text(slide_item["title"])
        h_p.font.size = Pt(19)
        h_p.font.bold = True
        h_p.font.color.rgb = NAVY
        h_p.alignment = PP_ALIGN.LEFT

        sub_p = h_tf.add_paragraph()
        sub_p.text = "JAGAT GURU NANAK DEV PUNJAB STATE OPEN UNIVERSITY, PATIALA | M.SC. (CS)"
        sub_p.font.size = Pt(9)
        sub_p.font.bold = True
        sub_p.font.color.rgb = MUTED_SLATE
        sub_p.alignment = PP_ALIGN.LEFT

        # Main Content Card (Crisp White Container)
        c_card = slide.shapes.add_shape(1, Inches(0.8), Inches(1.7), Inches(11.733), Inches(5.1))
        c_card.fill.solid()
        c_card.fill.fore_color.rgb = WHITE
        c_card.line.color.rgb = RGBColor(226, 232, 240)

        c_tf = c_card.text_frame
        c_tf.word_wrap = True
        c_tf.vertical_anchor = MSO_ANCHOR.TOP

        # Add bullets
        bullets = [clean_text(b) for b in slide_item["bullets"] if clean_text(b)]
        if bullets:
            for b_idx, bullet in enumerate(bullets):
                p = c_tf.paragraphs[0] if b_idx == 0 else c_tf.add_paragraph()
                p.text = "•   " + bullet
                p.font.size = Pt(13)
                p.font.color.rgb = SLATE
                p.space_after = Pt(10)
                p.level = 0
        else:
            p = c_tf.paragraphs[0]
            p.text = "•   " + clean_text(slide_item.get("raw", ""))
            p.font.size = Pt(13)
            p.font.color.rgb = SLATE

        # Bottom Footer on Slide
        f_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        f_tf = f_box.text_frame
        f_p = f_tf.paragraphs[0]
        f_p.text = f"Slide {idx} of {len(slides_data)}   |   Candidate: Sumit Singh (Roll No: PSOU/MSCS/2026/042)   |   JGND PSOU, Patiala"
        f_p.font.size = Pt(9)
        f_p.font.color.rgb = MUTED_SLATE

        # Presenter Notes (Spoken Speech Script & Timing)
        if slide_item.get("speech"):
            notes_tf = slide.notes_slide.notes_text_frame
            notes_tf.text = clean_text(slide_item["speech"])

    output_path = os.path.join(DOWNLOADS_DIR, file_name)
    prs.save(output_path)
    print(f"Successfully generated: {output_path} ({len(slides_data)} slides)")
    return output_path

def main():
    data_js_path = os.path.join(APP_DIR, "data.js")
    with open(data_js_path, "r", encoding="utf-8") as f:
        raw = f.read().split("window.ACADEMIC_DATA = ")[1].rstrip(";\n")
        data = json.loads(raw)

    # 1. Generate Major Project PPTX
    if "project_slides" in data and "parsed_slides" in data["project_slides"]:
        create_presentation(
            data["project_slides"]["parsed_slides"],
            "Major_Project_Presentation_MSCS-4-03P.pptx",
            is_major_project=True
        )

    # 2. Generate Seminar PPTX
    if "seminar_slides" in data and "parsed_slides" in data["seminar_slides"]:
        create_presentation(
            data["seminar_slides"]["parsed_slides"],
            "Seminar_Presentation_MSCS-4-04.pptx",
            is_major_project=False
        )

if __name__ == "__main__":
    main()
